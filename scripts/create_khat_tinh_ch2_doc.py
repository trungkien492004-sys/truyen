import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_chapter_2_docx(output_path, img_path):
    doc = docx.Document()

    # --- Page Setup (Book Margins: A4, 2.5cm top/bottom, 2.8cm left/right) ---
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.1)
    section.right_margin = Inches(1.1)

    # --- Header & Footer ---
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Khát Tình  |  Chương 02: Hố Sâu")
    hrun.font.name = "Palatino Linotype"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(157, 68, 110) # Dusty Rose

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("— ❦ —")
    frun.font.name = "Palatino Linotype"
    frun.font.size = Pt(9)
    frun.font.color.rgb = RGBColor(180, 120, 140)

    # --- Color Palette (Màu Ái Muội: Burgundy / Wine / Plum / Rose Gold / Soft Charcoal) ---
    COLOR_WINE_DEEP   = RGBColor(84, 18, 36)    # #541224 - Deep Burgundy
    COLOR_BURGUNDY    = RGBColor(107, 29, 47)   # #6B1D2F - Rich Bordeaux
    COLOR_ROSE_DUST   = RGBColor(157, 68, 110)  # #9D446E - Dusty Mauve
    COLOR_ROSE_ACCENT = RGBColor(186, 112, 137) # #BA7089 - Soft Rose Gold
    COLOR_BODY_TEXT   = RGBColor(35, 31, 32)    # #231F20 - Soft Charcoal Espresso
    COLOR_FLASHBACK   = RGBColor(85, 30, 48)    # #551E30 - Sensual Plum for Flashback
    
    HEX_BG_BOX        = "FDF7F9" # Ultra-soft pearl rose tint
    HEX_BORDER_ROSE   = "C98CA7" # Rose gold border
    HEX_DEEP_WINE     = "6B1D2F" # Deep wine

    # Set normal style font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Palatino Linotype'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_BODY_TEXT

    # --- Helper Functions ---
    def set_cell_shading(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_box_border(cell, color_hex, size_eighth_pts=10):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
                <w:left w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
                <w:bottom w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
                <w:right w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)

    def add_ornament_divider(symbol="❦  ✦  ❦"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(f"─── {symbol} ───")
        run.font.name = "Palatino Linotype"
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_ROSE_ACCENT

    def add_body_paragraph(text, is_first=False, italic=False, color=COLOR_BODY_TEXT):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(0)
        
        if is_first:
            p.paragraph_format.first_line_indent = Inches(0)
        else:
            p.paragraph_format.first_line_indent = Inches(0.28)

        run = p.add_run(text)
        run.font.name = "Palatino Linotype"
        run.font.size = Pt(11.5)
        run.font.italic = italic
        run.font.color.rgb = color
        return p

    # ==================== CHAPTER TITLE BLOCK ====================
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(20)
    p_top.paragraph_format.space_after = Pt(4)
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_series = p_top.add_run("K H Á T   T Ì N H")
    run_series.font.name = "Palatino Linotype"
    run_series.font.size = Pt(10.5)
    run_series.font.bold = True
    run_series.font.color.rgb = COLOR_ROSE_DUST

    p_mini = doc.add_paragraph()
    p_mini.paragraph_format.space_before = Pt(2)
    p_mini.paragraph_format.space_after = Pt(10)
    p_mini.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_mini = p_mini.add_run("· ❦ ·")
    r_mini.font.color.rgb = COLOR_ROSE_ACCENT
    r_mini.font.size = Pt(10)

    p_num = doc.add_paragraph()
    p_num.paragraph_format.space_before = Pt(6)
    p_num.paragraph_format.space_after = Pt(4)
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_num = p_num.add_run("CHƯƠNG 02")
    run_num.font.name = "Palatino Linotype"
    run_num.font.size = Pt(13)
    run_num.font.bold = True
    run_num.font.color.rgb = COLOR_BURGUNDY

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("HỐ SÂU")
    run_title.font.name = "Palatino Linotype"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_WINE_DEEP

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(24)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("❖  · · · ───────── ❦ ───────── · · ·  ❖")
    r_div.font.name = "Palatino Linotype"
    r_div.font.size = Pt(10)
    r_div.font.color.rgb = COLOR_ROSE_DUST

    # ==================== SCENE 1: PHÒNG THAY ĐỒ VỚI CHỊ NGÂN ====================
    add_body_paragraph(
        "Yên đứng trong góc khuất của phòng thay đồ riêng, ánh đèn vàng ấm phủ lên làn da trần như mật ong chảy. Cô vừa cởi bỏ chiếc váy ngoài, để lộ bộ nội y ren đen mỏng tang — bộ đồ mà Ngân đã chọn cho cô, bảo rằng \"đêm nay em phải là con đĩ đẹp nhất.\"",
        is_first=True
    )
    add_body_paragraph("\"Em đẹp quá...\"")
    add_body_paragraph("Giọng nói nhẹ nhàng vang lên sau lưng. Yên quay lại, thấy Ngân đang dựa vào khung cửa, ánh mắt chị ấy đắm đuối nhìn cô từ đầu đến chân. Ngân mặc một chiếc váy đỏ ôm sát, hai bầu vú căng tròn như muốn trào ra.")
    add_body_paragraph("\"Chị Ngân...\" Yên khẽ gọi, tay vô thức che ngực.")
    add_body_paragraph("Ngân bước đến, tay chị nhẹ nhàng gạt tay Yên xuống. \"Che làm gì? Đêm nay là đêm của những cơn thèm. Em không cần che giấu gì cả.\"")
    add_body_paragraph("Ngón tay chị chạm vào vai Yên, trượt dọc theo xương quai xanh, rồi dừng lại ở mép áo ngực. \"Em biết không? Chị nhìn em thế này... chị ướt hết cả.\"")
    add_body_paragraph("Yên nuốt nước bọt. Không khí trong phòng đột nhiên nóng bừng. \"Chị đừng...\"")
    add_body_paragraph("\"Đừng gì?\" Ngân cười khẩy, tay chị luồn vào sau lưng Yên, mở khóa áo ngực. \"Đêm nay, trước khi em bước vào thế giới của đàn ông... để chị dạy em cách thỏa mãn một người phụ nữ.\"")
    add_body_paragraph("Chiếc áo ngực rơi xuống. Hai bầu vú Yên lộ ra, đầu vú săn cứng vì kích thích. Ngân thở dốc, chị cúi xuống, miệng ngậm lấy một bên vú Yên.")
    add_body_paragraph("\"Ưm... ahh...\" Yên rên lên, tay bấu chặt vào vai Ngân.")
    add_body_paragraph("Lưỡi Ngân liếm láp quanh quầng vú, rồi mút mạnh đầu vú vào miệng, khiến Yên run lẩy bẩy. \"Chị... chị ơi...\"")
    add_body_paragraph("\"Im đi,\" Ngân thì thầm, tay kia xoa bóp bên vú còn lại. \"Để chị cho em biết thế nào là sung sướng thật sự. Không phải cái thứ giả tạo với thằng Khải yếu sinh lý kia.\"")
    add_body_paragraph("Yên muốn phản kháng, nhưng cơ thể cô phản bội. Khi Ngân đẩy cô ngồi xuống ghế, quỳ giữa hai chân cô, thì cô chỉ biết nằm im, thở dốc.")
    add_body_paragraph("Ngân kéo quần lót Yên xuống. Cái lồn đã ướt nhẹp lộ ra, mép ngoài hồng hào, hột le căng cứng như hạt đậu. \"Trời ơi... em ướt thế này mà bảo không muốn?\"")
    add_body_paragraph("\"Chị... em... ahh!\" Yên hét lên khi Ngân liếm một đường từ dưới lên trên, đầu lưỡi chạm vào hột le.")
    add_body_paragraph("Ngân bú liếm say sưa, tiếng \"chụp chụp\" vang lên trong phòng kín. Lưỡi chị đâm sâu vào lỗ lồn, rồi liếm lên liếm xuống, khiến Yên co giật từng hồi. \"Ngon không em? Có sướng không?\"")
    add_body_paragraph("\"Sướng... sướng quá chị ơi... ahh... đừng dừng...\" Yên rên rỉ, tay bấm chặt vào đầu Ngân, ép mặt chị sâu hơn vào háng mình.")
    add_body_paragraph("Nhưng rồi... cái cảm giác sung sướng đột ngột dừng lại. Ngân ngẩng lên, môi chị lấp lánh nước nhờn của Yên. \"Em thấy chưa? Đây mới là cách làm tình. Không phải cái thứ chó chết của thằng Khải.\"")
    add_body_paragraph("Yên thở hổn hển, lồn vẫn co thắt đòi hỏi. \"Sao... sao chị dừng...\"")
    add_body_paragraph("\"Vì em cần thằng cặc thật,\" Ngân cười, lau môi. \"Chứ lưỡi chị chỉ là món khai vị. Ra ngoài đi. Có một thằng đang chờ em. Thằng mà sẽ địt em như chưa từng được địt.\"")

    add_ornament_divider("❦  ✦  ❦")

    # ==================== SCENE 2: FLASHBACK ====================
    add_body_paragraph("Yên bước ra khỏi phòng thay đồ, váy đỏ ôm sát như lớp da thứ hai. Cô nhìn thấy Khải ở góc phòng, đang ve vãn một cô gái trẻ, tay hắn đặt trên đùi cô ta.", is_first=True)
    add_body_paragraph("Và đúng lúc đó... ký ức ùa về.")

    # Flashback blockquote / stylized section
    p_fb = doc.add_paragraph()
    p_fb.paragraph_format.space_before = Pt(8)
    p_fb.paragraph_format.space_after = Pt(6)
    p_fb.paragraph_format.first_line_indent = Inches(0)
    r_fb = p_fb.add_run("❖ Flashback — Ba năm trước, lần đầu tiên...")
    r_fb.font.name = "Palatino Linotype"
    r_fb.font.size = Pt(11)
    r_fb.font.bold = True
    r_fb.font.color.rgb = COLOR_ROSE_DUST

    add_body_paragraph("Căn phòng khách sạn sang trọng, ánh nến lung linh. Khải khi đó còn trẻ, còn sung. Anh ta đè Yên xuống giường, miệng gầm gừ như thú dữ.", italic=True, color=COLOR_FLASHBACK)
    add_body_paragraph("\"Em đẹp quá... để anh ăn em...\"", italic=True, color=COLOR_FLASHBACK)
    add_body_paragraph("Đó là lần đầu tiên Yên được bú lồn. Khải liếm từ gốc đến ngọn, lưỡi anh ta đánh vào hột le điên cuồng, khiến Yên rên rỉ tê dại. \"Ahh... anh ơi... sướng... sướng quá...\"", italic=True, color=COLOR_FLASHBACK)
    add_body_paragraph("\"Rên lớn lên,\" Khải ra lệnh, tay anh ta thọc hai ngón tay vào lỗ lồn mút chặt. \"Anh muốn nghe em rên.\"", italic=True, color=COLOR_FLASHBACK)
    add_body_paragraph("Yên rên như điên, cơ thể co giật. Đó là lần đầu tiên cô lên đỉnh, nước lồn phun ra ướt đẫm miệng Khải. Anh ta ngừng liếm, cười đắc ý, rồi leo lên người cô.", italic=True, color=COLOR_FLASHBACK)
    add_body_paragraph("Con cặc anh ta to, cứng, nóng bỏng. Khi anh ta đâm vào, Yên hét lên vì sung sướng. \"Ahh... địt em đi... địt mạnh vào...\"", italic=True, color=COLOR_FLASHBACK)
    add_body_paragraph("Khải địt như máy, tiếng \"bạch bạch\" vang lên khi đùi anh ta đập vào mông cô. Mỗi cú đâm đều chạm đến tử cung, khiến Yên khóc thét vì sướng. \"Địt nữa... địt mạnh nữa... em sắp ra...\"", italic=True, color=COLOR_FLASHBACK)
    add_body_paragraph("\"Ra đi em... ra cho anh...\"", italic=True, color=COLOR_FLASHBACK)
    add_body_paragraph("Yên lên đỉnh lần thứ hai, co thắt dữ dội. Khải cũng rên lớn, bắn tinh đầy lồn cô, tinh trùng nóng hổi chảy ra ngoài...", italic=True, color=COLOR_FLASHBACK)

    p_fb_end = doc.add_paragraph()
    p_fb_end.paragraph_format.space_before = Pt(8)
    p_fb_end.paragraph_format.space_after = Pt(6)
    p_fb_end.paragraph_format.first_line_indent = Inches(0)
    r_fb_end = p_fb_end.add_run("❖ Hiện tại...")
    r_fb_end.font.name = "Palatino Linotype"
    r_fb_end.font.size = Pt(11)
    r_fb_end.font.bold = True
    r_fb_end.font.color.rgb = COLOR_ROSE_DUST

    add_body_paragraph("Yên chớp mắt, cảm thấy lồn mình ướt nhẹp chỉ vì nhớ lại. Nhưng rồi cô nhìn Khải bây giờ — gã đàn ông đang cười nói với cô gái kia, dương vật chắc chắn đang mềm oặt trong quần.")
    add_body_paragraph("Lần thứ hai, thứ ba... rồi hàng trăm lần sau đó, Khải dần yếu đi. Không còn những cú địt mạnh bạo, không còn những lần bú liếm say sưa. Thay vào đó là sự bất lực, là con cặc mềm nhũn không thể cứng, là những đêm Yên nằm đó với lồn ướt át mà không được thỏa mãn.")
    add_body_paragraph("Những thèm khát của cái lồn dâm điếm không được thỏa mãn khiến cô như chìm dần vào hố sâu tuyệt vọng. Những chờ đợi mong mỏi tan biến dần theo năm tháng, biến cô thành một cái xác biết đi với cái lồn khô héo.")

    add_ornament_divider("❦  ✦  ❦")

    # ==================== SCENE 3: GẶP GỠ DƯƠNG ====================
    add_body_paragraph("\"Em ổn không?\"", is_first=True)
    add_body_paragraph("Giọng nói trầm ấm vang lên bên tai. Yên quay lại, thấy một người đàn ông đeo mặt nạ đen, chỉ lộ đôi mắt sắc bén như săn mồi.")
    add_body_paragraph("Anh ta cao to, mặc vest đen ôm sát, để lộ cơ ngực săn chắc. Và dù không nhìn thấy, Yên có thể cảm nhận — bên dưới lớp quần tây ấy, có một con cặc đang cứng ngắc chờ đợi.")
    add_body_paragraph("\"Em trông như kẻ đói,\" anh ta cười, giọng đầy ẩn ý. \"Đói cặc phải không?\"")
    add_body_paragraph("Yên nuốt nước bọt, lồn cô co thắt. \"Anh là ai?\"")
    add_body_paragraph("\"Người sẽ cho em ăn,\" anh ta đáp, tay nắm lấy eo cô kéo sát. \"Ăn no căng. Ăn đến khi em khóc thét vì sung sướng.\"")
    add_body_paragraph("Yên run lẩy bẩy. Cô biết mình nên chạy trốn. Nhưng cái lồn thèm khát của cô đã quyết định thay cô.")
    add_body_paragraph("\"Được,\" cô thì thầm, tay vô thức chạm vào ngực anh ta. \"Cho em ăn đi.\"")
    add_body_paragraph("Người đàn ông cười lớn, ôm cô vào lòng. \"Gọi anh là Dương. Và đêm nay, em sẽ quên mẹ thằng Khải yếu sinh lý kia là ai.\"")
    add_body_paragraph("Hắn hôn cô. Không phải nụ hôn dịu dàng, mà là nụ hôn của kẻ chiếm hữu — môi hắn cắn mạnh môi cô, lưỡi đâm sâu vào miệng cô như con cặc sắp đâm vào lồn cô.")
    add_body_paragraph("Yên rên rỉ trong cổ họng, cơ thể cô cháy bừng.")
    add_body_paragraph("Hố sâu tuyệt vọng cuối cùng cũng có người kéo cô lên.")
    add_body_paragraph("Bằng con cặc cứng ngắc của hắn.")

    # ==================== EMBEDDED SEXY ILLUSTRATION ====================
    table_ill = doc.add_table(rows=1, cols=1)
    table_ill.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_ill.autofit = False
    table_ill.columns[0].width = Inches(5.6)

    cell_ill = table_ill.cell(0, 0)
    set_cell_shading(cell_ill, HEX_BG_BOX)
    set_cell_box_border(cell_ill, HEX_BORDER_ROSE, size_eighth_pts=12)

    # Caption Tag
    cp1 = cell_ill.paragraphs[0]
    cp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp1.paragraph_format.space_before = Pt(8)
    cp1.paragraph_format.space_after = Pt(6)
    r_ill_tag = cp1.add_run("✦  M I N H   H Ọ A   C H Ư Ơ N G   2  ✦")
    r_ill_tag.font.name = "Palatino Linotype"
    r_ill_tag.font.size = Pt(9.5)
    r_ill_tag.font.bold = True
    r_ill_tag.font.color.rgb = COLOR_BURGUNDY

    # Image
    cp_img = cell_ill.add_paragraph()
    cp_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_img.paragraph_format.space_before = Pt(2)
    cp_img.paragraph_format.space_after = Pt(6)
    run_img = cp_img.add_run()
    run_img.add_picture(img_path, width=Inches(4.8))

    # Caption text
    cp2 = cell_ill.add_paragraph()
    cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp2.paragraph_format.space_before = Pt(4)
    cp2.paragraph_format.space_after = Pt(8)
    r_ill_text = cp2.add_run("“Bên bức tường đá lạnh, nhốt giữa hơi thở nóng rực và dục vọng kẻ săn mồi”")
    r_ill_text.font.name = "Palatino Linotype"
    r_ill_text.font.size = Pt(11)
    r_ill_text.font.italic = True
    r_ill_text.font.bold = True
    r_ill_text.font.color.rgb = COLOR_WINE_DEEP

    # ==================== SCENE 4: ÉP VÀO TƯỜNG & ĐỐI THOẠI NÓNG BỎNG ====================
    add_body_paragraph("Yên lùi lại, lưng chạm vào tường lạnh. Dương tiến sát, hai tay anh ta chống lên tường hai bên đầu cô, nhốt cô trong không gian chật hẹp của mùi nước hoa nam tính và hơi thở nóng rực.", is_first=True)
    add_body_paragraph("\"Chạy à?\" Dương cười khẩy, đầu lưỡi liếm nhẹ môi. \"Vừa rồi em rên trong miệng anh nghe sướng lắm mà?\"")
    add_body_paragraph("\"Im đi...\" Yên quay mặt đi, nhưng má cô đỏ bừng.")
    add_body_paragraph("Dương cúi xuống, môi áp sát vành tai cô, thổi nhẹ: \"Cái lồn em ướt rồi đúng không? Anh ngửi thấy mùi dâm thủy của em đây... thèm cặc lắm rồi à?\"")
    add_body_paragraph("\"Anh... anh đừng nói bậy...\" Yên run lẩy bẩy, nhưng hai đầu vú cô đã cứng ngắc chọc vào ngực anh ta.")
    add_body_paragraph("Dương tay trượt xuống, ngón tay chạm nhẹ vào eo cô rồi dừng ở mép váy. \"Anh đoán... bên dưới cái váy này, cái quần lót em đang ướt nhẹp. Em có muốn anh kiểm tra không? Hay em tự cởi cho anh xem?\"")
    add_body_paragraph("\"Không...\" Yên lắc đầu, nhưng tay cô lại nắm chặt áo sơ mi của anh ta.")
    add_body_paragraph("\"Nói dối,\" Dương cắn nhẹ vành tai cô, khiến Yên rên nhẹ. \"Em biết anh nhìn thấy gì không? Em nhìn anh như con đĩ đói cặc. Đôi mắt em đang van xin anh địt em đấy.\"")
    add_body_paragraph("Yên nuốt nước bọt, giọng run run: \"Anh... anh là ai mà tự tin thế?\"")
    add_body_paragraph("\"Anh là người sẽ cho em biết thế nào là đàn ông thật sự,\" Dương thì thầm, tay anh ta nâng cằm cô lên bắt nhìn vào mắt mình. \"Không phải thằng Khải yếu sinh lý kia. Anh sẽ địt em đến khi em khóc thét, đến khi em quên cả tên mình. Em muốn không?\"")
    add_body_paragraph("Yên im lặng, cơ thể cô run rẩy. Cô muốn — trời ơi cô muốn quá — nhưng nỗi sợ hãi và lý trí vẫn còn.")
    add_body_paragraph("\"Trả lời anh,\" Dương ra lệnh, tay anh ta siết nhẹ cổ cô. \"Em có muốn anh đút cặc vào lồn em không? Có muốn anh địt em như con đĩ không?\"")
    add_body_paragraph("\"Ưm...\" Yên rên nhẹ, đầu gối mềm nhũn. \"Em... em...\"")
    add_body_paragraph("\"Nói đi,\" Dương trêu chọc, đùi anh ta chèn vào giữa hai chân cô, cọ xát vào cái lồn đang ướt qua lớp váy. \"Nói anh nghe xem. Em muốn gì?\"")
    add_body_paragraph("Yên hít thở gấp gáp, cảm giác cứng rắn của đùi anh ta chạm vào nơi nhạy cảm khiến cô muốn gào thét. \"Em... em muốn...\"")
    add_body_paragraph("\"Muốn gì?\"")
    add_body_paragraph("\"Muốn... ahh...\" Yên chưa kịp nói hết, Dương đã hôn cô ngấu nghiến.")
    add_body_paragraph("Nụ hôn như lửa đốt. Lưỡi anh ta đâm sâu, quấn lấy lưỡi cô, mút chặt đến nghẹt thở. Tay anh ta xoa nắn ngực cô qua váy, ngón tay se se đầu vú khiến Yên rên rỉ trong cổ họng.")
    add_body_paragraph("\"Ngon không?\" Dương thả môi ra, hơi thở dồn dập. \"Môi anh ngon hơn thằng Khải chứ?\"")
    add_body_paragraph("\"Ngon... ngon hơn nhiều...\" Yên thừa nhận, mặt đỏ bừng.")
    add_body_paragraph("Dương cười đắc ý, tay anh ta trượt xuống dưới váy cô, chạm vào đùi trong mềm mại. \"Anh đoán... giờ anh sờ vào quần lót em, sẽ thấy ướt nhẹp. Đúng không?\"")
    add_body_paragraph("\"Không... đừng...\" Yên giật mình, nhưng không đẩy anh ta ra.")
    add_body_paragraph("\"Nói dối,\" Dương thì thầm, tay anh ta tiến sát mép quần lót. \"Em biết anh thích gì không? Anh thích nhìn con đĩ đói cặc vừa thèm vừa sợ. Em đang sợ anh địt em đến chết à?\"")
    add_body_paragraph("Yên gật nhẹ, mắt ướt át. \"Em sợ... sợ anh...\"")
    add_body_paragraph("\"Sợ anh cái gì?\"")
    add_body_paragraph("\"Sợ... sợ anh làm em nghiện...\" Yên thở dốc thừa nhận. \"Sợ anh địt em xong... em không thể sống thiếu cặc anh...\"")
    add_body_paragraph("Dương mắt sáng rực, anh ta cười lớn: \"Thông minh đấy. Anh sẽ làm em nghiện. Nghiện đến mức đêm nào cũng nằm dạng chân chờ anh địt. Em tin không?\"")
    add_body_paragraph("Yên không trả lời. Cô sợ — sợ chính mình, sợ cơn thèm khát đang nuốt chửng lý trí.")
    add_body_paragraph("\"Thôi được,\" Dương bỗng rút tay ra, lùi lại một bước. \"Em chưa sẵn sàng. Anh không thích địt con đĩ còn do dự.\"")
    add_body_paragraph("Yên ngơ ngác nhìn anh ta, cơ thể cô trống rỗng, lồn cô co thắt đòi hỏi. \"Anh...\"")
    add_body_paragraph("\"Đi đi,\" Dương quay lưng, giọng lạnh băng. \"Về với thằng Khải yếu sinh lý của em đi. Để cái lồn dâm đãng của em khô héo tiếp đi.\"")
    add_body_paragraph("Yên cắn môi, nước mắt chực trào. Cô xấu hổ vì bị từ chối, xấu hổ vì chính sự thèm khát của mình. Cô quay người, chạy vụt ra khỏi phòng.")
    add_body_paragraph("Cánh cửa đóng sầm lại.")
    add_body_paragraph("Dương đứng đó, con cặc trong quần cứng ngắc đau đớn. Anh ta thở hổn hển, tức giận vì bị cự tuyệt, vì cái lồn ngon lành đã bị tuột mất.")
    add_body_paragraph("\"Được thôi...\" anh ta gầm gừ, lấy điện thoại ra nhắn tin. \"Nếu em không chịu... thì anh sẽ địt người khác. Và em sẽ hối hận.\"")
    add_body_paragraph("Tin nhắn gửi đi: \"Vào phòng anh ngay. Cả hai đứa.\"")

    # ==================== END BLOCK ====================
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(24)
    p_end.paragraph_format.space_after = Pt(16)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("❖   HẾT CHƯƠNG 2   ❖")
    r_end.font.name = "Palatino Linotype"
    r_end.font.size = Pt(12.5)
    r_end.font.bold = True
    r_end.font.color.rgb = COLOR_BURGUNDY

    p_final = doc.add_paragraph()
    p_final.paragraph_format.space_before = Pt(4)
    p_final.paragraph_format.space_after = Pt(10)
    p_final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_final = p_final.add_run("❦  ❦  ❦")
    r_final.font.color.rgb = COLOR_ROSE_ACCENT
    r_final.font.size = Pt(11)

    doc.save(output_path)
    print(f"Document saved successfully at: {output_path}")

if __name__ == "__main__":
    img = r"C:\Users\Kien\.gemini\antigravity\brain\c7856a80-f02d-4922-97d3-e2d2e1081fc5\khat_tinh_chapter_2_1793107077569.jpg"
    out = r"C:\Users\Kien\Downloads\Khat_Tinh_Chuong_02.docx"
    create_chapter_2_docx(out, img)
