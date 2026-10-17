import os
import sys
import gzip
import urllib.request
import re
import html
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

def clean_html(raw_html):
    if not raw_html:
        return []
    text = re.sub(r'<br\s*/?>', '\n', raw_html, flags=re.IGNORECASE)
    text = re.sub(r'</p>', '\n\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<[^>]+>', '', text)
    text = html.unescape(text)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', text)
    paras = [line.strip() for line in text.split('\n') if line.strip()]
    return paras

def main():
    headers = {'User-Agent': 'Mozilla/5.0'}
    req = urllib.request.Request('https://www.wattpad.com/api/v3/stories/273327250', headers=headers)
    with urllib.request.urlopen(req) as resp:
        import json
        story_data = json.loads(resp.read().decode('utf-8'))

    parts = story_data.get('parts', [])
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Bìa
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("ĐIỆN THOẠI THẦN KỲ Ở TU CHÂN GIỚI")
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.name = 'Times New Roman'
    r.font.color.rgb = RGBColor(15, 23, 42)

    p_a = doc.add_paragraph()
    p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_a.paragraph_format.space_after = Pt(20)
    r_a = p_a.add_run("Tác giả: Niken / Kin Kên")
    r_a.font.size = Pt(14)
    r_a.font.italic = True
    r_a.font.name = 'Times New Roman'
    r_a.font.color.rgb = RGBColor(100, 116, 139)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run(f"Bản trích xuất và ghép nối toàn bộ {len(parts)} phần từ Wattpad (story/273327250)")
    r_sub.font.size = Pt(11)
    r_sub.font.name = 'Times New Roman'
    r_sub.font.color.rgb = RGBColor(148, 163, 184)

    doc.add_page_break()

    for idx, p_info in enumerate(parts):
        pid = p_info.get('id')
        p_title = p_info.get('title')
        if p_title.isdigit():
            full_title = f"Phần {idx + 1}: Chương {p_title}"
        else:
            full_title = f"Phần {idx + 1}: {p_title}"

        req_text = urllib.request.Request(f"https://www.wattpad.com/apiv2/storytext?id={pid}", headers=headers)
        with urllib.request.urlopen(req_text) as r_text:
            raw = r_text.read()
            try:
                content = gzip.decompress(raw).decode('utf-8', errors='ignore')
            except Exception:
                content = raw.decode('utf-8', errors='ignore')

        h = doc.add_paragraph()
        h_run = h.add_run(full_title)
        h_run.font.size = Pt(16)
        h_run.font.bold = True
        h_run.font.name = 'Times New Roman'
        h_run.font.color.rgb = RGBColor(180, 83, 9)
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(10)

        paras = clean_html(content)
        for p_text in paras:
            para = doc.add_paragraph()
            para.paragraph_format.line_spacing = 1.25
            para.paragraph_format.space_after = Pt(4)
            para.paragraph_format.first_line_indent = Inches(0.3)
            r_body = para.add_run(p_text)
            r_body.font.name = 'Times New Roman'
            r_body.font.size = Pt(12)
            r_body.font.color.rgb = RGBColor(30, 41, 59)

        if idx < len(parts) - 1:
            doc.add_page_break()

    out_docx = r'C:\Users\Kien\Downloads\Dien_Thoai_Than_Ky_Wattpad_Toan_Tap.docx'
    doc.save(out_docx)
    print(f"✅ Đã tạo thêm file DOCX: {out_docx}")

if __name__ == '__main__':
    main()
