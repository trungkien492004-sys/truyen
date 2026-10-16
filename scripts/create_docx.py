import json
import re
import html
import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def clean_html_to_paragraphs(raw_html):
    if not raw_html:
        return []
    text = re.sub(r'<br\s*/?>', '\n', raw_html, flags=re.IGNORECASE)
    text = re.sub(r'</p>', '\n\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<[^>]+>', '', text)
    text = html.unescape(text)
    # Strip XML invalid control characters
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', text)
    paras = []
    for line in text.split('\n'):
        line_clean = line.strip()
        if line_clean:
            paras.append(line_clean)
    return paras

def main():
    # Fetch from Supabase
    import subprocess
    print("Đang lấy toàn bộ dữ liệu 79 chương từ database...")
    subprocess.run(["node", "scripts/fetch_story_7_for_word.js"], check=True)

    with open('temp_story_7.json', 'r', encoding='utf-8') as f:
        data = json.load(f)

    story = data['story']
    chapters = data['chapters']
    print(f"Đã đọc {len(chapters)} chương từ database.")

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Trang tiêu đề / Trang bìa
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(12)
    run = title_p.add_run(story.get('title', 'Điện Thoại Thần Kỳ Ở Tu Chân Giới'))
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(24, 30, 45)

    author_p = doc.add_paragraph()
    author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_p.paragraph_format.space_after = Pt(24)
    author_run = author_p.add_run(f"Tác giả: {story.get('author', 'Niken')}")
    author_run.font.size = Pt(14)
    author_run.font.italic = True
    author_run.font.name = 'Times New Roman'
    author_run.font.color.rgb = RGBColor(100, 116, 139)

    # Giới thiệu nội dung
    if story.get('description'):
        desc_box = doc.add_paragraph()
        desc_box.paragraph_format.space_before = Pt(18)
        desc_box.paragraph_format.space_after = Pt(6)
        d_run = desc_box.add_run("TÓM TẮT NỘI DUNG:")
        d_run.font.bold = True
        d_run.font.size = Pt(13)
        d_run.font.name = 'Times New Roman'
        
        desc_p = doc.add_paragraph(story.get('description'))
        desc_p.paragraph_format.line_spacing = 1.2
        desc_p.paragraph_format.space_after = Pt(24)
        for r in desc_p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)

    meta_p = doc.add_paragraph(f"Tổng số: {len(chapters)} chương | Đã biên tập và tinh chỉnh chuẩn theo yêu cầu tác giả")
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in meta_p.runs:
        r.font.italic = True
        r.font.size = Pt(10)
        r.font.name = 'Times New Roman'
        r.font.color.rgb = RGBColor(148, 163, 184)

    doc.add_page_break()

    print(f"Đang ghi {len(chapters)} chương vào file Word...")

    for idx, ch in enumerate(chapters):
        ch_num = ch.get('chapter_number', idx + 1)
        ch_title = ch.get('title') or f"Chương {ch_num}"
        ch_title = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', ch_title)

        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h_run = h.add_run(ch_title)
        h_run.font.size = Pt(16)
        h_run.font.bold = True
        h_run.font.name = 'Times New Roman'
        h_run.font.color.rgb = RGBColor(217, 119, 6)
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(10)

        paras = clean_html_to_paragraphs(ch.get('content', ''))
        for p_text in paras:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.first_line_indent = Inches(0.3)
            r = p.add_run(p_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(30, 41, 59)

        if idx < len(chapters) - 1:
            doc.add_page_break()

    out_file1 = r'C:\Users\Kien\Downloads\Dien_Thoai_Than_Ky_O_Tu_Chan_Gioi.docx'
    out_file2 = r'C:\Users\Kien\Downloads\Dien_Thoai_Than_Ky_O_Tu_Chan_Gioi_Full_79_Chuong.docx'
    
    doc.save(out_file1)
    doc.save(out_file2)
    
    size_mb = os.path.getsize(out_file1) / (1024 * 1024)
    print(f"✅ THÀNH CÔNG! Đã lưu file Word {len(chapters)} chương hoàn chỉnh vào:\n- {out_file1}\n- {out_file2}")
    print(f"Dung lượng: {size_mb:.2f} MB")

    if os.path.exists('temp_story_7.json'):
        os.remove('temp_story_7.json')

if __name__ == '__main__':
    main()
