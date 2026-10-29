import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_chapter_5_docx(output_path, img_path=None):
    doc = docx.Document()

    # --- Page Setup (Book Margins: A4, 2.5cm top/bottom, 2.8cm left/right) ---
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.1)
    section.right_margin = Inches(1.1)

    # --- Header & Footer ---
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Khát Tình  |  Chương 05: Bốn Người Một Phòng")
    hrun.font.name = "Palatino Linotype"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(157, 68, 110) # Dusty Rose

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("— ❦ —")
    frun.font.name = "Palatino Linotype"
    frun.font.size = Pt(9)
    frun.font.color.rgb = RGBColor(180, 120, 140)

    # --- Color Palette ---
    COLOR_WINE_DEEP   = RGBColor(84, 18, 36)    # #541224 - Deep Burgundy
    COLOR_BURGUNDY    = RGBColor(107, 29, 47)   # #6B1D2F - Rich Bordeaux
    COLOR_ROSE_DUST   = RGBColor(157, 68, 110)  # #9D446E - Dusty Mauve
    COLOR_ROSE_ACCENT = RGBColor(186, 112, 137) # #BA7089 - Soft Rose Gold
    COLOR_BODY_TEXT   = RGBColor(35, 31, 32)    # #231F20 - Soft Charcoal Espresso
    
    HEX_BG_BOX        = "FDF7F9" # Ultra-soft pearl rose tint
    HEX_BORDER_ROSE   = "C98CA7" # Rose gold border
    HEX_DEEP_WINE     = "6B1D2F" # Deep wine

    # Set normal style font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Palatino Linotype'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_BODY_TEXT

    # --- Helper Functions ---
    def set_cell_shading(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_box_border(cell, color_hex, size_eighth_pts=10):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
                <w:left w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
                <w:bottom w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
                <w:right w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)

    def add_ornament_divider(symbol="❦  ✦  ❦"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(f"─── {symbol} ───")
        run.font.name = "Palatino Linotype"
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_ROSE_ACCENT

    def add_body_paragraph(text, is_first=False, italic=False, color=COLOR_BODY_TEXT):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(0)
        
        if is_first:
            p.paragraph_format.first_line_indent = Inches(0)
        else:
            p.paragraph_format.first_line_indent = Inches(0.28)

        run = p.add_run(text)
        run.font.name = "Palatino Linotype"
        run.font.size = Pt(11.5)
        run.font.italic = italic
        run.font.color.rgb = color
        return p

    # ==================== CHAPTER TITLE BLOCK ====================
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(20)
    p_top.paragraph_format.space_after = Pt(4)
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_series = p_top.add_run("K H Á T   T Ì N H")
    run_series.font.name = "Palatino Linotype"
    run_series.font.size = Pt(10.5)
    run_series.font.bold = True
    run_series.font.color.rgb = COLOR_ROSE_DUST

    p_mini = doc.add_paragraph()
    p_mini.paragraph_format.space_before = Pt(2)
    p_mini.paragraph_format.space_after = Pt(10)
    p_mini.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_mini = p_mini.add_run("· ❦ ·")
    r_mini.font.color.rgb = COLOR_ROSE_ACCENT
    r_mini.font.size = Pt(10)

    p_num = doc.add_paragraph()
    p_num.paragraph_format.space_before = Pt(6)
    p_num.paragraph_format.space_after = Pt(4)
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_num = p_num.add_run("CHƯƠNG 05")
    run_num.font.name = "Palatino Linotype"
    run_num.font.size = Pt(13)
    run_num.font.bold = True
    run_num.font.color.rgb = COLOR_BURGUNDY

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("BỐN NGƯỜI MỘT PHÒNG")
    run_title.font.name = "Palatino Linotype"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_WINE_DEEP

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(24)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("❖  · · · ───────── ❦ ───────── · · ·  ❖")
    r_div.font.name = "Palatino Linotype"
    r_div.font.size = Pt(10)
    r_div.font.color.rgb = COLOR_ROSE_DUST

    # ==================== SCENE 1: LỜI MỜI BẤT NGỜ ====================
    add_body_paragraph(
        "Một tuần sau những chuỗi ngày vật lộn trong cơn ngứa ngáy thể xác và tội lỗi giày vò...",
        is_first=True
    )
    add_body_paragraph(
        "Yên đang cặm cụi rửa bát trong gian bếp nhỏ thì Khải bước về nhà, vẻ mặt lộ rõ nét hào hứng khác thường:"
    )
    add_body_paragraph(
        "\"Cuối tuần này có một bữa tiệc kín tại biệt thự riêng trên đồi. Em chuẩn bị rồi đi cùng anh nhé? Chỉ có vài người bạn thân thiết thôi.\""
    )
    add_body_paragraph(
        "Đôi bàn tay đang cầm chiếc bát sứ của Yên bỗng khựng lại giữa chừng. Cô hiểu — trong thâm tâm cô linh cảm rõ mồn một đó là loại tiệc tùng gì. Nhưng cô không có quyền từ chối. Khải từng yêu thương mẹ cô tha thiết, đã nuôi nấng, che chở cho cô suốt ba năm ròng rã. Cô nợ gã cả một ân tình lớn hơn trời biển."
    )
    add_body_paragraph(
        "\"Dạ... em sẽ đi cùng anh,\" cô cúi đầu thì thào, giọng nói khẽ run lên."
    )

    add_ornament_divider("❦  ✦  ❦")

    # ==================== SCENE 2: BIỆT THỰ KHU ĐỒI & CUỘC ĐỐI ĐẦU BỐN NGƯỜI ====================
    add_body_paragraph(
        "Đêm thứ Bảy — Biệt thự sang trọng tọa lạc cô tịch trên sườn đồi lộng gió.",
        is_first=True
    )
    add_body_paragraph(
        "Yên khoác lên mình chiếc váy dạ hội đỏ rực bó sát mà Khải đã cẩn thận chọn lựa, phần cổ áo khoét sâu táo bạo phơi bày trọn vẹn rãnh ngực căng tràn mời gọi. Từng bước chân cô tiến vào căn phòng VIP thượng hạng, lồng ngực đập liên hồi như muốn nổ tung vì căng thẳng tột độ."
    )
    add_body_paragraph(
        "Cánh cửa gỗ mở toang. Và rồi... Yên đứng chết trân tại chỗ."
    )
    add_body_paragraph(
        "Dương đang đứng sừng sững ở đó, một tay tự nhiên ôm lấy vòng eo mảnh khảnh của một cô gái trẻ vô cùng xinh đẹp trong chiếc váy trắng tinh khôi. Cô gái ấy nép sát vào lòng hắn như một chú chim non e ấp tựa cành cây, đôi mắt to tròn ươn ướt long lanh, toát lên vẻ nhu mì, yếu đuối đến nao lòng."
    )
    add_body_paragraph(
        "\"Để anh giới thiệu,\" Dương khẽ nhếch mép cười, ánh mắt sắc như dao găm xoáy thẳng vào Yên. \"Đây là Linh, vị hôn thê của anh. Gia đình hai bên ép anh phải cưới cô ấy vào tháng sau.\""
    )
    add_body_paragraph(
        "Linh ngoan ngoãn cúi đầu chào, giọng nói nhỏ nhẹ, yếu ớt như tiếng gió thoảng: \"Em chào chị Yên... Em thường nghe anh Dương nhắc về chị...\""
    )
    add_body_paragraph(
        "Bàn tay Yên siết chặt lại, móng tay găm sâu vào lòng bàn tay đau nhói. Một cơn ghen tuông trào dâng mãnh liệt — trời đất ơi, cô ghen đến mức phát điên, ghen đến mức muốn lao vào cào xé cô gái ngây thơ kia. Nhưng đồng thời, sâu thẳm trong lòng cô lại trào dâng một nỗi xót xa, cay đắng. Linh nhìn Dương bằng ánh mắt tôn sùng, mù quáng tuyệt đối, cam tâm tình nguyện làm một con rối mặc hắn giật dây."
    )
    add_body_paragraph(
        "Đôi mắt Khải bỗng sáng rực lên khi quét từ đầu đến chân Linh — cô gái này đúng là gu sở thích bệnh hoạn của gã: non nớt, dễ bảo, ngực nở nang và làn da trắng ngần nõn nà. Khải liền lên tiếng đề nghị với giọng điệu nhớp nháp:"
    )
    add_body_paragraph(
        "\"Hay là đêm nay chúng ta chơi trò gì đó mới lạ một chút đi? Đổi đối tác nhé? Tôi chơi Linh, còn cậu... cứ tự nhiên với Yên.\""
    )
    add_body_paragraph(
        "Linh run rẩy cả người, đôi mắt đẫm lệ nhìn sang Dương như muốn van nài được che chở. Nhưng đáp lại cô chỉ là một cái gật đầu lạnh lùng đến tàn nhẫn của hắn:"
    )
    add_body_paragraph(
        "\"Được thôi. Đêm nay Linh sẽ chiều chuộng anh Khải tất cả mọi thứ. Cô ấy ngoan ngoãn và nhu nhược lắm, bảo làm gì cũng sẽ nghe theo.\""
    )
    add_body_paragraph(
        "\"Quỳ xuống,\" Dương lạnh giọng ra lệnh cho Linh."
    )
    add_body_paragraph(
        "Linh không dám chần chừ nửa giây, lập tức quỳ sụp xuống, vạt váy trắng phủ mềm mại trên tấm thảm nhung đen tuyền, đầu cúi thật thấp, hai bàn tay đặt ngay ngắn trên đùi. Cô ngoan ngoãn như một con mèo đã được thuần hóa kỹ lưỡng, chỉ biết câm lặng chờ đợi chủ nhân định đoạt."
    )
    add_body_paragraph(
        "Yên bàng hoàng lùi lại một bước, tấm lưng chạm vào bức tường lạnh buốt: \"Anh... rốt cuộc các anh định làm cái quái gì vậy?\""
    )
    add_body_paragraph(
        "Dương sải bước tiến lại gần, bàn tay thô bạo túm lấy mép váy đỏ của Yên giật mạnh một đường — tiếng vải rách toạc xoẹt một tiếng chói tai trong không gian tĩnh mịch."
    )
    add_body_paragraph(
        "\"Đêm nay,\" hắn cúi sát, phả hơi thở nóng rực đầy mùi rượu và nam tính vào vành tai cô, \"cả bốn người chúng ta sẽ ở đây. Cùng trong một căn phòng. Cùng tận mắt nhìn nhau làm tình. Và anh... sẽ địt nát lồn em ngay trước mặt chồng chưa cưới của em.\""
    )

    add_ornament_divider("❦  ✦  ❦")

    # ==================== SCENE 3: BẢN GIAO HƯỞNG SẮC DỤC CỰC HẠN ====================
    add_body_paragraph(
        "Khải lúc này đã bị dục vọng thiêu đốt, thân hình gầy gò của gã bước tới đứng sừng sững trước mặt Linh. Gã kéo phéc-mơ-tuya quần tây, cây cặc ngắn ngủn, mềm nhũn teo tóp của gã lập tức lộ ra ngoài.",
        is_first=True
    )
    add_body_paragraph(
        "\"Ngậm vào đi em... Mau bú cho nó cứng lên...\" Khải thô lỗ thúc giục."
    )
    add_body_paragraph(
        "Linh ngước nhìn Dương thêm một lần cuối cùng trong tuyệt vọng, nhưng chỉ nhận lại được ánh mắt thờ ơ, lạnh lùng như băng tảng. Cô khẽ cắn môi, há rộng khuôn miệng nhỏ nhắn ngậm trọn lấy thanh thịt èo uột của Khải, mút mát nhẹ nhàng từng chút một."
    )
    add_body_paragraph(
        "\"Ưm... chùn chụt...\" Tiếng bú mút dâm đãng bắt đầu vang lên róc rách trong căn phòng kín."
    )
    add_body_paragraph(
        "Cùng lúc đó, Dương lôi Yên ném mạnh lên chiếc giường tròn cỡ lớn đặt ngay giữa phòng — nơi có tấm gương tráng bạc khổng lồ gắn trên trần nhà phản chiếu lại từng ngóc ngách của cuộc truy hoan. Hắn đè ngửa cô ra nệm, một tay xé toạc phần còn lại của chiếc váy dạ hội."
    )
    add_body_paragraph(
        "\"Nhìn lên trần đi,\" Dương ghì cằm cô, chỉ tay lên tấm gương lớn phía trên. \"Nhìn xem người đàn ông của em đang được vị hôn thê ngoan ngoãn của anh quỳ gối bú cặc thế nào. Em có thấy ghen không?\""
    )
    add_body_paragraph(
        "Yên ngước mắt nhìn lên — qua tấm kính trong suốt, cô thấy rõ mồn một cảnh Khải đang đê mê vuốt ve mái tóc của Linh, khuôn mặt gã giãn ra hưởng thụ tột cùng khi cô gái trẻ quỳ gối phục vụ dưới háng mình. Cảnh tượng cuckold nhục nhã ấy vừa khiến lòng cô quặn thắt tủi hổ, nhưng kỳ lạ thay, lại kích thích bản năng dâm đãng nơi hạ bộ đến điên cuồng — dâm thủy tuôn trào ướt nhẹp cả đáy quần lót ren."
    )
    add_body_paragraph(
        "\"Anh... anh đừng làm thế...\" Yên thở dốc nghẹn ngào, nhưng đôi chân thon dài lại bất giác mở rộng ra theo từng chuyển động của Dương khi hắn lột phăng chiếc quần lót mỏng manh của cô xuống."
    )
    add_body_paragraph(
        "Dương cởi bỏ quần áo, con cặc khổng lồ lập tức nảy bật ra ngoài, đỏ ửng, gân guốc cuộn tròn nổi rõ mồn một như chuẩn bị nổ tung. Hắn đè trọn thân hình vạm vỡ lên người Yên, quy đầu to bè nóng rực dí sát vào đóa hoa xuân đang ướt sũng dâm thủy."
    )
    add_body_paragraph(
        "\"Em đã mong chờ giây phút này bao lâu rồi?\" Dương thì thầm, bàn tay to lớn xoa nắn thô bạo bầu ngực căng mẩy của cô. \"Đêm nào cũng tự thủ dâm rồi nức nở gọi tên anh... có phải không?\""
    )
    add_body_paragraph(
        "Yên đỏ bừng mặt, xấu hổ đến mức chỉ muốn độn thổ — hắn biết tất cả, không một bí mật nào của cô qua mắt được hắn."
    )
    add_body_paragraph(
        "\"Ưm... phải... em thừa nhận... đêm nào em cũng tự móc lồn mà nhớ đến anh...\" Yên bật khóc thừa nhận trong nhục nhã ê chề."
    )
    add_body_paragraph(
        "Dương nhếch mép cười đắc ý, rồi bất thình lình gồng hông thúc mạnh một cú thấu trời — cả cây côn thịt hộ pháp đâm lút cán vào sâu trong âm đạo chật hẹp, chạm thẳng đến tận cổ tử cung Yên."
    )
    add_body_paragraph(
        "\"Ahhh!\" Yên thét lên một tiếng thất thanh, hai bàn tay bấu chặt vào bờ vai rắn chắc của Dương, móng tay cào rách da thịt hắn. \"To quá! Cặc anh to dã man! Địt em đi! Thúc mạnh nữa vào anh ơi!\""
    )

    add_ornament_divider("❦  ✦  ❦")

    # ==================== SCENE 4: TRUY HOAN TỘI LỖI TRƯỚC TẤM GƯƠNG TRẦN ====================
    add_body_paragraph(
        "Ở góc phòng bên kia, Khải bất giác quay đầu lại nhìn. Gã tận mắt chứng kiến người vợ sắp cưới dịu dàng, kín đáo của mình đang bị một người đàn ông khác cuồng bạo giã nát trên giường, khuôn mặt Yên đỏ ửng vì khoái lạc tột đỉnh, bờ môi quyến rũ không ngừng rên rỉ van nài.",
        is_first=True
    )
    add_body_paragraph(
        "Nhưng thay vì bốc hỏa cuồng nộ, cảnh tượng ngoại tình công khai và trụy lạc tột bậc ấy lại đánh thức cơn dục vọng biến thái sâu kín trong con người gã — thanh cặc teo tóp của Khải cuối cùng cũng cương cứng thật sự trong vòm họng của Linh."
    )
    add_body_paragraph(
        "\"Liếm hai hòn dái của anh đi em...\" Khải thô bạo vỗ bồm bộp vào đầu Linh như đang thuần phục một con chó cái."
    )
    add_body_paragraph(
        "Linh ngoan ngoãn cúi thấp xuống, đưa đầu lưỡi ướt át quấn lấy hai hòn ngọc hoàn nhỏ thó của Khải, ngậm chặt từng hòn mà bú liếm. Những giọt nước mắt uất ức lặng lẽ rỉ ra nơi khóe mi, nhưng cô vẫn cắn răng phục tùng — bởi lời dặn của Dương là mệnh lệnh tuyệt đối mà cô không bao giờ dám trái ý."
    )
    add_body_paragraph(
        "\"Nhìn chồng em đi,\" Dương thúc từng cú dập long trời lở đất vào trong lồn Yên, mỗi cú va chạm đều khiến thành giường gỗ rung chuyển bần bật. \"Gã ta đang sung sướng tận hưởng vị hôn thê của anh kìa. Còn em... em đang bị anh địt như một con đĩ dâm đãng ngay trước mắt gã. Có sướng không?\""
    )
    add_body_paragraph(
        "\"Sướng! Sướng phát điên lên được anh ơi!\" Yên mất hết liêm sỉ, hai chân quặp chặt lấy hông Dương, lắc hông đón nhận từng cú nhấp thô bạo. \"Địt chết em đi! Sâu quá... Cặc anh đâm sâu đến nát tử cung em rồi!\""
    )
    add_body_paragraph(
        "Dương bất ngờ đổi tư thế — hắn kéo bổng Yên ngồi dậy, để cô ngồi lọt thỏm trên cây cặc to lớn của mình, mặt đối mặt sát rạt."
    )
    add_body_paragraph(
        "\"Nhìn thẳng vào mắt anh. Cấm em được nhìn sang gã chồng bất lực của em. Đêm nay, toàn bộ linh hồn và thể xác em đều thuộc về anh.\""
    )
    add_body_paragraph(
        "Yên ôm chặt lấy cổ Dương, nhún nhảy điên cuồng trên khúc thịt gân guốc, hai bầu ngực trần nảy lên từng nhịp rạo rực. Lúc này cô đã hoàn toàn quên đi sự hiện diện của Khải, quên sạch mọi rào cản luân lý tội lỗi — trong đầu cô chỉ còn lại khoái cảm nhục dục tột đỉnh đang cuộn trào như núi lửa."
    )
    add_body_paragraph(
        "Phía bên kia, Khải đã đủ cương cứng để giao hợp. Gã thô bạo đẩy Linh nằm sấp xuống tấm thảm nhung, cô gái nhu nhược chỉ biết nhắm nghiền mắt, cam chịu dạng rộng hai chân đón nhận: \"Em... em sẵn sàng rồi...\""
    )
    add_body_paragraph(
        "Khải lập tức đâm phập thanh thịt nhỏ thó vào kẽ lồn đã ướt đẫm nước miếng của Linh. \"Ưm... Khít quá... Chặt chẽ quá...\" Khải thở dốc rên rỉ, nhấp từng nhịp chậm chạp từ phía sau."
    )
    add_body_paragraph(
        "Linh chỉ biết cắn chặt môi khẽ rên từng tiếng yếu ớt, không dám kêu la lớn vì sợ Dương chán ghét. Cô nằm bất động, đôi bàn tay mảnh dẻ bấu chặt vào sợi thảm nhung đen, nhẫn nhục chịu đựng từng cú dập của người đàn ông già nua."
    )
    add_body_paragraph(
        "\"Nhìn họ kìa,\" Dương ghé môi cắn nhẹ vào cổ Yên, bàn tay bóp nghiến bầu ngực mọng. \"Nhìn vị hôn thê của anh đang bị chồng em đè ra địt. Và nhìn chính em xem... đang cởi truồng cưỡi cặc anh ngay trước mắt gã.\""
    )
    add_body_paragraph(
        "Yên ngửa cổ nhìn lên trần nhà — tấm gương lớn phản chiếu trọn vẹn cảnh tượng hoang dâm trụy lạc đến mức không tưởng: bốn thân xác trần truồng không mảnh vải che thân, hai cặp đôi đang điên cuồng quấn lấy nhau, tiếng da thịt va đập chan chát cùng tiếng rên xiết hòa quyện thành một bản hòa ca tội lỗi rợn người."
    )
    add_body_paragraph(
        "Cô thấy Khải đang dập liên hồi vào mông Linh, mặt gã đỏ gay vì thỏa mãn. Và cô thấy chính mình — đang trần như nhộng lắc hông điên dại trên cây cặc khổng lồ của Dương, hai bầu vú rung bần bật, hơi thở dồn dập trong cơn say tình."
    )
    add_body_paragraph(
        "\"Em ra... Em ra mất rồi anh ơi... Ahhh!\" Yên co thắt dữ dội, dâm thủy phun trào xối xả tưới đẫm gốc cặc Dương."
    )
    add_body_paragraph(
        "\"Anh cũng bắn đây!\" Dương gầm lên một tiếng vang động cả căn phòng, dập mạnh cú cuối cùng rồi phóng từng đợt tinh dịch đặc quánh, nóng rực sâu vào tận đáy tử cung Yên."
    )
    add_body_paragraph(
        "Cùng lúc đó, Khải cũng hét lên một tiếng thỏa mãn: \"Ra rồi... Ra rồi... Ahh!\" Gã xuất tinh vào sâu bên trong Linh, con cặc teo tóp giật giật vài cái rồi mềm oặt tuột ra ngoài thảm."
    )

    add_ornament_divider("❦  ✦  ❦")

    # ==================== SCENE 5: TÀN CUỘC HOAN LẠC & SỰ PHỤC TÙNG VĨNH VIỄN ====================
    add_body_paragraph(
        "Tàn cuộc hoan lạc.",
        is_first=True
    )
    add_body_paragraph(
        "Bốn con người nằm rã rời trong căn phòng ngột ngạt, không khí nồng nặc mùi mồ hôi, dâm thủy và tinh dịch tanh nồng."
    )
    add_body_paragraph(
        "Linh nằm co ro một góc trên thảm, kẽ bướm vẫn rỉ ra dòng tinh dịch trắng đục của Khải, hai hàng nước mắt lăn dài nhưng bờ môi vẫn mấp máy thì thào đầy sợ sệt: \"Em... em đã làm tốt đúng không anh Dương?\""
    )
    add_body_paragraph(
        "Dương chẳng buồn liếc nhìn cô lấy một cái. Hắn ôm chặt Yên vào lồng ngực vạm vỡ ướt đẫm mồ hôi, vuốt ve bờ vai mềm mại của cô."
    )
    add_body_paragraph(
        "\"Kể từ đêm nay,\" hắn thì thầm với giọng điệu chiếm hữu tuyệt đối, \"em chính thức thuộc về anh. Hoàn toàn và mãi mãi.\""
    )
    add_body_paragraph(
        "Yên nằm im trong vòng tay hắn, bụng dưới vẫn ấm nóng dòng tinh dịch sôi sục của Dương. Cô biết bản thân đã phản bội Khải một cách nhục nhã và triệt để nhất. Thế nhưng lạ lùng thay, trong lòng cô không hề có lấy một tia oán hận hay tiếc nuối. Thứ duy nhất còn sót lại... chỉ là khoái lạc ngút ngàn đan xen cùng nỗi trống rỗng vô tận."
    )

    # ==================== EMBEDDED ILLUSTRATION (IF PROVIDED) ====================
    if img_path and os.path.exists(img_path):
        table_ill = doc.add_table(rows=1, cols=1)
        table_ill.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_ill.autofit = False
        table_ill.columns[0].width = Inches(5.6)

        cell_ill = table_ill.cell(0, 0)
        set_cell_shading(cell_ill, HEX_BG_BOX)
        set_cell_box_border(cell_ill, HEX_BORDER_ROSE, size_eighth_pts=12)

        # Caption Tag
        cp1 = cell_ill.paragraphs[0]
        cp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp1.paragraph_format.space_before = Pt(8)
        cp1.paragraph_format.space_after = Pt(6)
        r_ill_tag = cp1.add_run("✦  M I N H   H Ọ A   C H Ư Ơ N G   5  ✦")
        r_ill_tag.font.name = "Palatino Linotype"
        r_ill_tag.font.size = Pt(9.5)
        r_ill_tag.font.bold = True
        r_ill_tag.font.color.rgb = COLOR_BURGUNDY

        # Image
        cp_img = cell_ill.add_paragraph()
        cp_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp_img.paragraph_format.space_before = Pt(2)
        cp_img.paragraph_format.space_after = Pt(6)
        run_img = cp_img.add_run()
        run_img.add_picture(img_path, width=Inches(4.8))

        # Caption text
        cp2 = cell_ill.add_paragraph()
        cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp2.paragraph_format.space_before = Pt(4)
        cp2.paragraph_format.space_after = Pt(8)
        r_ill_text = cp2.add_run("“Bốn thân xác cuốn vào trò chơi đổi bạn tình đồi bại, phản chiếu dưới tấm gương trần của dục vọng không đáy”")
        r_ill_text.font.name = "Palatino Linotype"
        r_ill_text.font.size = Pt(11)
        r_ill_text.font.italic = True
        r_ill_text.font.bold = True
        r_ill_text.font.color.rgb = COLOR_WINE_DEEP

    # ==================== CHAPTER QUOTE BOX ====================
    table_q = doc.add_table(rows=1, cols=1)
    table_q.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_q.autofit = False
    table_q.columns[0].width = Inches(5.8)

    cell_q = table_q.cell(0, 0)
    set_cell_shading(cell_q, HEX_BG_BOX)
    set_cell_box_border(cell_q, HEX_BORDER_ROSE, size_eighth_pts=12)

    # Caption Tag
    cp1_q = cell_q.paragraphs[0]
    cp1_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp1_q.paragraph_format.space_before = Pt(10)
    cp1_q.paragraph_format.space_after = Pt(6)
    r_q_tag = cp1_q.add_run("✦  Đ I Ể M   N H Ấ N   C H Ư Ơ N G   5  ✦")
    r_q_tag.font.name = "Palatino Linotype"
    r_q_tag.font.size = Pt(9.5)
    r_q_tag.font.bold = True
    r_q_tag.font.color.rgb = COLOR_BURGUNDY

    # Caption text
    cp2_q = cell_q.add_paragraph()
    cp2_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp2_q.paragraph_format.space_before = Pt(4)
    cp2_q.paragraph_format.space_after = Pt(12)
    cp2_q.paragraph_format.left_indent = Inches(0.2)
    cp2_q.paragraph_format.right_indent = Inches(0.2)
    r_q_text = cp2_q.add_run("“Trước tấm gương trần phản chiếu bản giao hưởng xác thịt đồi bại, mọi luân thường đạo lý đều tan thành mây khói — chỉ còn lại sự phục tùng tuyệt đối trước dục vọng của kẻ săn mồi.”")
    r_q_text.font.name = "Palatino Linotype"
    r_q_text.font.size = Pt(11)
    r_q_text.font.italic = True
    r_q_text.font.bold = True
    r_q_text.font.color.rgb = COLOR_WINE_DEEP

    # ==================== END BLOCK ====================
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(24)
    p_end.paragraph_format.space_after = Pt(16)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("❖   HẾT CHƯƠNG 5   ❖")
    r_end.font.name = "Palatino Linotype"
    r_end.font.size = Pt(12.5)
    r_end.font.bold = True
    r_end.font.color.rgb = COLOR_BURGUNDY

    p_final = doc.add_paragraph()
    p_final.paragraph_format.space_before = Pt(4)
    p_final.paragraph_format.space_after = Pt(10)
    p_final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_final = p_final.add_run("❦  ❦  ❦")
    r_final.font.color.rgb = COLOR_ROSE_ACCENT
    r_final.font.size = Pt(11)

    doc.save(output_path)
    print(f"Document saved successfully at: {output_path}")

if __name__ == "__main__":
    img = r"C:\Users\Kien\.gemini\antigravity\scratch\truyen\public\images\illustrations\khat-tinh-ch5.jpg"
    out_download = r"C:\Users\Kien\Downloads\Khat_Tinh_Chuong_05.docx"
    out_scratch = r"C:\Users\Kien\.gemini\antigravity\scratch\truyen\Khat_Tinh_Chuong_05.docx"
    
    create_chapter_5_docx(out_download, img)
    create_chapter_5_docx(out_scratch, img)
