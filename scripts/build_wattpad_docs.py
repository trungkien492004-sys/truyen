import json
import re
import html
import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def clean_html(raw_html):
    if not raw_html:
        return []
    text = re.sub(r'<br\s*/?>', '\n', raw_html, flags=re.IGNORECASE)
    text = re.sub(r'</p>', '\n\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<[^>]+>', '', text)
    text = html.unescape(text)
    # Strip XML incompatible control characters
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', text)
    paras = []
    for line in text.split('\n'):
        line_clean = line.strip()
        if line_clean:
            paras.append(line_clean)
    return paras

def create_doc_from_parts(title_story, author, parts_list, out_path, subtitle="Bản trích xuất từ Wattpad"):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Bìa
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(title_story)
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.name = 'Times New Roman'
    r.font.color.rgb = RGBColor(24, 30, 45)

    p_a = doc.add_paragraph()
    p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_a.paragraph_format.space_after = Pt(24)
    r_a = p_a.add_run(f"Tác giả: {author}")
    r_a.font.size = Pt(13)
    r_a.font.italic = True
    r_a.font.name = 'Times New Roman'
    r_a.font.color.rgb = RGBColor(100, 116, 139)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run(f"{subtitle} ({len(parts_list)} phần/chương)")
    r_sub.font.size = Pt(11)
    r_sub.font.name = 'Times New Roman'
    r_sub.font.color.rgb = RGBColor(148, 163, 184)

    doc.add_page_break()

    for idx, pt in enumerate(parts_list):
        pt_title = pt.get('title') or f"Phần {idx + 1}"
        pt_title = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', pt_title)
        
        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h_run = h.add_run(pt_title)
        h_run.font.size = Pt(16)
        h_run.font.bold = True
        h_run.font.name = 'Times New Roman'
        h_run.font.color.rgb = RGBColor(217, 119, 6)
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(10)

        paras = clean_html(pt.get('content', ''))
        for p_text in paras:
            para = doc.add_paragraph()
            para.paragraph_format.line_spacing = 1.25
            para.paragraph_format.space_after = Pt(4)
            para.paragraph_format.first_line_indent = Inches(0.3)
            r_body = para.add_run(p_text)
            r_body.font.name = 'Times New Roman'
            r_body.font.size = Pt(12)
            r_body.font.color.rgb = RGBColor(30, 41, 59)

        if idx < len(parts_list) - 1:
            doc.add_page_break()

    doc.save(out_path)
    size_mb = os.path.getsize(out_path) / (1024 * 1024)
    print(f"✅ Đã lưu: {out_path} ({size_mb:.2f} MB)")

def main():
    # 1. Đọc dữ liệu Wattpad
    with open('wattpad_parts_raw.json', 'r', encoding='utf-8') as f:
        wp_data = json.load(f)

    wp_story = wp_data['story']
    wp_parts = wp_data['parts']

    # File 1: Toàn bộ 10 phần trên Wattpad
    path_wp = r'C:\Users\Kien\Downloads\Dien_Thoai_Than_Ky_Wattpad_10_Parts.docx'
    create_doc_from_parts(
        title_story="Điện Thoại Thần Kỳ Ở Tu Chân Giới",
        author=wp_story.get('user', {}).get('name', 'Kin Kên'),
        parts_list=wp_parts,
        out_path=path_wp,
        subtitle="Toàn bộ 10 phần đăng trên Wattpad (story/273327250)"
    )

    # File 2: Chỉ các chương tiếp theo trên Wattpad (Chương 84, 95, 108)
    next_parts = [
        {'title': 'Chương 84: ' + wp_parts[7]['title'], 'content': wp_parts[7]['content']},
        {'title': 'Chương 95: ' + wp_parts[8]['title'], 'content': wp_parts[8]['content']},
        {'title': 'Chương 108: ' + wp_parts[9]['title'], 'content': wp_parts[9]['content']}
    ]

    path_next_only = r'C:\Users\Kien\Downloads\Dien_Thoai_Than_Ky_Cac_Chuong_Tiep_Theo_Wattpad.docx'
    create_doc_from_parts(
        title_story="Điện Thoại Thần Kỳ Ở Tu Chân Giới",
        author="Niken / Kin Kên",
        parts_list=next_parts,
        out_path=path_next_only,
        subtitle="Các chương tiếp theo từ nguồn Wattpad (Chương 84, 95, 108)"
    )

    # File 3: Full gộp 77 chương gốc + 3 chương tiếp theo từ Wattpad
    with open('temp_story_7.json', 'r', encoding='utf-8') as f:
        db_data = json.load(f)
    
    db_chapters = db_data['chapters']
    combined_chapters = []
    for c in db_chapters:
        combined_chapters.append({
            'title': c.get('title') or f"Chương {c.get('chapter_number')}",
            'content': c.get('content', '')
        })
    combined_chapters.extend(next_parts)

    path_full_combined = r'C:\Users\Kien\Downloads\Dien_Thoai_Than_Ky_Full_Kem_Wattpad.docx'
    create_doc_from_parts(
        title_story="Điện Thoại Thần Kỳ Ở Tu Chân Giới (Full Bản Mới Nhất)",
        author="Niken / Kin Kên",
        parts_list=combined_chapters,
        out_path=path_full_combined,
        subtitle=f"Bản tổng hợp đầy đủ nhất (77 chương web + 3 chương tiếp theo từ Wattpad)"
    )

    # Dọn dẹp file tạm
    if os.path.exists('wattpad_parts_raw.json'):
        os.remove('wattpad_parts_raw.json')
    if os.path.exists('temp_story_7.json'):
        os.remove('temp_story_7.json')

if __name__ == '__main__':
    main()
