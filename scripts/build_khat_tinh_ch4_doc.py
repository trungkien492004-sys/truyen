import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_chapter_4_docx(output_path, img_path=None):
    doc = docx.Document()

    # --- Page Setup (Book Margins: A4, 2.5cm top/bottom, 2.8cm left/right) ---
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.1)
    section.right_margin = Inches(1.1)

    # --- Header & Footer ---
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Khát Tình  |  Chương 04: Tội Lỗi Và Cơn Ngứa")
    hrun.font.name = "Palatino Linotype"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(157, 68, 110) # Dusty Rose

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("— ❦ —")
    frun.font.name = "Palatino Linotype"
    frun.font.size = Pt(9)
    frun.font.color.rgb = RGBColor(180, 120, 140)

    # --- Color Palette ---
    COLOR_WINE_DEEP   = RGBColor(84, 18, 36)    # #541224 - Deep Burgundy
    COLOR_BURGUNDY    = RGBColor(107, 29, 47)   # #6B1D2F - Rich Bordeaux
    COLOR_ROSE_DUST   = RGBColor(157, 68, 110)  # #9D446E - Dusty Mauve
    COLOR_ROSE_ACCENT = RGBColor(186, 112, 137) # #BA7089 - Soft Rose Gold
    COLOR_BODY_TEXT   = RGBColor(35, 31, 32)    # #231F20 - Soft Charcoal Espresso
    COLOR_FLASHBACK   = RGBColor(85, 30, 48)    # #551E30 - Sensual Plum for Flashback
    
    HEX_BG_BOX        = "FDF7F9" # Ultra-soft pearl rose tint
    HEX_BORDER_ROSE   = "C98CA7" # Rose gold border
    HEX_DEEP_WINE     = "6B1D2F" # Deep wine

    # Set normal style font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Palatino Linotype'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_BODY_TEXT

    # --- Helper Functions ---
    def set_cell_shading(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_box_border(cell, color_hex, size_eighth_pts=10):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
                <w:left w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
                <w:bottom w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
                <w:right w:val="single" w:sz="{size_eighth_pts}" w:space="0" w:color="{color_hex}"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)

    def add_ornament_divider(symbol="❦  ✦  ❦"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(f"─── {symbol} ───")
        run.font.name = "Palatino Linotype"
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_ROSE_ACCENT

    def add_body_paragraph(text, is_first=False, italic=False, color=COLOR_BODY_TEXT):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(0)
        
        if is_first:
            p.paragraph_format.first_line_indent = Inches(0)
        else:
            p.paragraph_format.first_line_indent = Inches(0.28)

        run = p.add_run(text)
        run.font.name = "Palatino Linotype"
        run.font.size = Pt(11.5)
        run.font.italic = italic
        run.font.color.rgb = color
        return p

    # ==================== CHAPTER TITLE BLOCK ====================
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(20)
    p_top.paragraph_format.space_after = Pt(4)
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_series = p_top.add_run("K H Á T   T Ì N H")
    run_series.font.name = "Palatino Linotype"
    run_series.font.size = Pt(10.5)
    run_series.font.bold = True
    run_series.font.color.rgb = COLOR_ROSE_DUST

    p_mini = doc.add_paragraph()
    p_mini.paragraph_format.space_before = Pt(2)
    p_mini.paragraph_format.space_after = Pt(10)
    p_mini.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_mini = p_mini.add_run("· ❦ ·")
    r_mini.font.color.rgb = COLOR_ROSE_ACCENT
    r_mini.font.size = Pt(10)

    p_num = doc.add_paragraph()
    p_num.paragraph_format.space_before = Pt(6)
    p_num.paragraph_format.space_after = Pt(4)
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_num = p_num.add_run("CHƯƠNG 04")
    run_num.font.name = "Palatino Linotype"
    run_num.font.size = Pt(13)
    run_num.font.bold = True
    run_num.font.color.rgb = COLOR_BURGUNDY

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("TỘI LỖI VÀ CƠN NGỨA")
    run_title.font.name = "Palatino Linotype"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_WINE_DEEP

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(24)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("❖  · · · ───────── ❦ ───────── · · ·  ❖")
    r_div.font.name = "Palatino Linotype"
    r_div.font.size = Pt(10)
    r_div.font.color.rgb = COLOR_ROSE_DUST

    # ==================== SCENE 1: CHUYẾN TAXI ĐÊM & NỖI CẮN RỨT ====================
    add_body_paragraph(
        "Yên ngồi cô độc ở băng ghế sau của chiếc taxi đang lao vút qua những con phố đêm vắng lặng. Đầu ngón tay cô run rẩy chạm nhẹ lên bờ môi mềm — nơi dường như vẫn còn vương vấn dư vị cuồng nhiệt, mùi thuốc lá thượng hạng và hơi thở nam tính nồng nàn đến nghẹt thở của Dương.",
        is_first=True
    )
    add_body_paragraph(
        "Một cơn rùng mình chạy dọc sống lưng khiến từng sợi lông tơ của cô dựng đứng. Không phải vì luồng khí lạnh từ điều hòa xe hơi, mà vì cơn sóng tội lỗi khủng khiếp đang cuộn trào, giày vò và cắn rứt lương tâm cô từng cơn."
    )
    add_body_paragraph(
        "\"Em đúng là đồ lăng loàn trắc nết... đồ đĩ thèm khát...\" Yên nghiến chặt răng tự nguyền rủa chính mình, hai hàng nước mắt nóng hổi lặng lẽ lăn dài trên gò má đỏ bừng vì xấu hổ."
    )
    add_body_paragraph(
        "Bàn tay run rẩy rút vội chiếc điện thoại ra khỏi túi xách. Cô dứt khoát chặn số điện thoại của Dương, xóa sạch mọi tin nhắn, xóa hết từng dấu vết vụng trộm vừa nhen nhóm. Cô tự dặn lòng phải chôn chặt đêm nay vào quên lãng vĩnh viễn, xem tất cả như một cơn ác mộng hoang đường, một ảo giác đồi bại nhất thời mà cô trót sa chân."
    )
    add_body_paragraph(
        "Về đến căn hộ quen thuộc, Khải đã ngủ say mèm trên ghế sofa tự bao giờ, tiếng thở nặng nề đều đặn vang lên trong phòng khách. Yên đứng lặng người ngắm nhìn gã — người đàn ông từng yêu thương mẹ cô tha thiết, người đã bảo bọc và mang lại cho cô chỗ dựa cuộc sống suốt những năm tháng khốn khó, cơ hàn nhất."
    )
    add_body_paragraph(
        "Dù Khải bất lực, yếu sinh lý, dù gã đạo đức giả và đầy rẫy thói hư tật xấu, nhưng với cô, gã vẫn là ân nhân lớn nhất cuộc đời."
    )
    add_body_paragraph(
        "\"Em xin lỗi anh...\" Yên quỳ gối bên mép sofa, thì thầm trong tiếng nấc nghẹn ngào rồi cúi xuống đặt một nụ hôn dịu dàng lên vầng trán nhăn nheo của Khải. \"Từ giờ em sẽ là một người vợ ngoan ngoãn... Em thề sẽ không bao giờ phản bội anh thêm một lần nào nữa.\""
    )

    add_ornament_divider("❦  ✦  ❦")

    # ==================== SCENE 2: NGƯỜI VỢ MẪU MỰC & CƠ THỂ PHẢN BỘI ====================
    add_body_paragraph(
        "Một tuần lễ sau đó — Yên dốc lòng biến mình thành một người vợ mẫu mực đến hoàn hảo.",
        is_first=True
    )
    add_body_paragraph(
        "Mỗi sớm mai, cô dậy thật sớm nấu bữa điểm tâm tinh tươm, ân cần chăm sóc và chiều chuộng Khải hết mực. Tối đến, cô chủ động xoa bóp vai gáy cho gã sau một ngày dài làm việc, thậm chí còn ngượng ngùng cởi phăng cúc áo ngực để mặc cho Khải tự do nắn bóp đôi gò bồng đảo căng tròn — điều mà suốt ba năm qua cô chưa từng một lần chủ động mở lời."
    )
    add_body_paragraph(
        "\"Hôm nay em làm sao thế?\" Khải nhướn mày ngạc nhiên, bàn tay thô kệch nhào nặn một bên vú cô. \"Nhìn em lạ lùng lắm đấy.\""
    )
    add_body_paragraph(
        "\"Em chỉ muốn làm cho anh được vui thôi mà,\" Yên mỉm cười hiền dịu, nhưng nơi sâu thẳm tâm can cô chỉ là một khoảng không trống rỗng, lạnh ngắt như hầm băng."
    )
    add_body_paragraph(
        "Đêm muộn, khi Khải đã chìm sâu vào giấc ngủ say và ngáy khò khè bên cạnh, Yên vẫn trằn trọc nằm trơ trọi, đôi mắt mở to nhìn trân trân lên trần nhà tối om."
    )
    add_body_paragraph(
        "Bàn tay của Khải vẫn đặt hờ hững trên bầu ngực trần của cô, nhưng xúc cảm thô vụng ấy chẳng thể mang lại cho cô lấy một tia kích thích nhỏ nhoi. Bất chợt, những ký ức nhục dục đêm đó lại ồ ạt tràn về như thác lũ cuốn phăng mọi rào giậu phòng thủ."
    )
    add_body_paragraph(
        "Khuôn mặt góc cạnh đầy ma mị của Dương... Đôi môi nóng bỏng như muốn ngấu nghiến lấy cô... Hơi thở nam tính phả vào vành tai rực lửa... Cây côn thịt khổng lồ, cứng ngắc như thanh sắt đè nghiến vào đùi non và kẽ lồn cô..."
    )
    add_body_paragraph(
        "\"Không... Không được nghĩ nữa...\" Yên khẽ lắc đầu, cố gắng xua tan ảo ảnh tội lỗi đang bóp nghẹt tâm trí."
    )
    add_body_paragraph(
        "Thế nhưng cơ thể cô lại hoàn toàn phản bội lý trí. Nơi sâu kín giữa hai đùi bỗng trào dâng một cảm giác ngứa ngáy, râm ran đến phát điên. Đáy quần lót mỏng manh đã ướt đẫm dâm thủy từ lúc nào."
    )
    add_body_paragraph(
        "Những ngón tay mảnh khảnh bất giác luồn sâu xuống mép quần lót ren — lần đầu tiên trong suốt hai mươi mấy năm cuộc đời, cô tự tay vuốt ve chính mình chỉ vì nhớ nhung và thèm khát cây cặc của một người đàn ông khác ngoài chồng sắp cưới."
    )
    add_body_paragraph(
        "\"Ưm... Dương... A...\" Yên khẽ thở hắt ra trong vô thức, đầu ngón tay ma sát liên hồi lên hạt le đang cương cứng giần giật."
    )
    add_body_paragraph(
        "Bất chợt sực tỉnh, cô giật bắn mình rút vội bàn tay nhớp nháp ra ngoài, lòng ngập tràn hoảng loạn: \"Không... Mày điên rồi Yên... Mày không được làm như thế...\""
    )

    add_ornament_divider("❦  ✦  ❦")

    # ==================== SCENE 3: NHỮNG CUỘC TAO NGỘ "TÌNH CỜ" ĐẦY QUỶ QUYỆT ====================
    add_body_paragraph(
        "Thế nhưng số phận — hay đúng hơn là sự sắp đặt tinh vi của kẻ đi săn mồi — không buông tha cho cô.",
        is_first=True
    )
    add_body_paragraph(
        "Ngày thứ ba sau đêm tiệc, Yên đi siêu thị mua sắm, bỗng bắt gặp Dương đang đứng ngay bên quầy rau củ tươi. Hắn đút một tay vào túi quần âu, nhìn cô bằng ánh mắt sắc lạnh như dao cạo, khóe môi khẽ nhếch lên nụ cười đầy trêu ngươi:"
    )
    add_body_paragraph(
        "\"Em cũng thích ăn cà chua à? Anh nhớ... em đặc biệt thích màu đỏ rực rỡ.\""
    )
    add_body_paragraph(
        "Yên hoảng hốt buông rơi cả giỏ đồ, vội vã quay đầu tháo chạy như kẻ trốn chạy bóng ma."
    )
    add_body_paragraph(
        "Ngày thứ năm, Yên đến trung tâm spa chăm sóc da để giải tỏa cơn căng thẳng thần kinh triền miên. Vừa bước chân vào khu xông hơi thư giãn, đập vào mắt cô lại là bóng hình quen thuộc đến nghẹt thở. Dương nằm dài trên chiếc giường nệm bên cạnh, chiếc khăn tắm trắng chỉ quấn hờ hững ngang hông, để lộ toàn bộ lồng ngực vạm vỡ, những múi cơ bụng rắn như tạc và làn da màu đồng nam tính bỏng rẫy."
    )
    add_body_paragraph(
        "\"Có muốn qua đây massage cho anh một chút không?\" Hắn cất giọng trầm khàn đầy ma lực quyến rũ."
    )
    add_body_paragraph(
        "Yên run rẩy cắn chặt môi đến suýt bật máu, thêm một lần nữa quay lưng bỏ chạy trối chết trong sự bấn loạn cực độ."
    )

    add_ornament_divider("❦  ✦  ❦")

    # ==================== SCENE 4: KHÁCH SẠN 5 SAO & CẢNH TƯỢNG CẤM KỴ ====================
    add_body_paragraph(
        "Đến ngày thứ bảy, Yên có chuyến công tác ngắn ngày tại một khách sạn 5 sao sang trọng bậc nhất trung tâm thành phố.",
        is_first=True
    )
    add_body_paragraph(
        "Tối muộn, cô ngồi ăn tối một mình tại nhà hàng lầu cao. Sau khi uống cạn hai ly vang đỏ đậm đà, men rượu ấm nồng bắt đầu ngấm vào từng mạch máu, khiến bước chân cô có chút lâng lâng, chuếnh choáng khi đứng dậy đi tìm nhà vệ sinh."
    )
    add_body_paragraph(
        "Trong cơn mơ màng, khi đầu óc vẫn ngập tràn những ảo ảnh bấn loạn về Dương, đôi bàn tay cô vô thức đẩy nhầm cánh cửa có đề biển kim loại sáng bóng: \"NAM\"."
    )
    add_body_paragraph(
        "Và rồi... Yên đứng chết trân tại chỗ."
    )
    add_body_paragraph(
        "Dương đang đứng sừng sững trước bệ tiểu bằng đá cẩm thạch trắng muốt, tấm lưng rộng lớn quay nửa về phía cửa ra vào. Hắn không hề mặc quần dài. Cây cặc thô dài khủng khiếp được bàn tay gân guốc nắm hờ, đang phóng ra một dòng nước tiểu vàng óng, tiếng dòng nước va đập róc rách vào bệ sứ vang vọng rõ mồn một trong không gian tĩnh mịch của căn phòng kín."
    )
    add_body_paragraph(
        "Yên trợn tròn mắt, hơi thở như ngưng trệ, hoàn toàn bất động không thể nào rời mắt khỏi cảnh tượng cấm kỵ trước mặt."
    )
    add_body_paragraph(
        "Cây cặc ấy — nó to lớn gấp bội phần thứ mềm nhẽo, bất lực của Khải, thô dài, gân máu tím thẫm cuồn cuộn nổi cộm như loài mãng xà săn mồi, đầu khấc to bè nở rộng tim tím vì máu tụ dồn ứ. Dòng nước tiểu mạnh mẽ, nóng rực phun xối xả rơi xuống lòng bệ tiểu trắng toát."
    )
    add_body_paragraph(
        "Yên nuốt ực một ngụm nước bọt, cổ họng khô khốc ran lên như thiêu như đốt. Trong một thoáng điên loạn của bản năng dục vọng, trong đầu cô bỗng lóe lên một ý nghĩ đồi bại và dâm đãng đến rợn người — cô tưởng tượng dòng nước tiểu nóng hổi, mặn chát ấy đang phóng thẳng vào khuôn miệng đang há rộng của mình, tràn xuống cuống họng, tưới đẫm lên hai bầu ngực trần... và cô sẽ điên cuồng nuốt ừng ực từng ngụm một như đang đón nhận thứ thần dược khoái lạc quý giá nhất trần đời..."
    )
    add_body_paragraph(
        "\"Thích nhìn đến thế cơ à?\""
    )
    add_body_paragraph(
        "Thanh âm trầm đục đầy tà khí đột ngột vang lên xé toạc ảo ảnh của Yên."
    )
    add_body_paragraph(
        "Dương đã thong thả xoay người lại. Bàn tay hắn vẫn lười biếng cầm lấy khúc thịt khổng lồ chưa kịp hạ nhiệt, vài giọt nước tiểu còn sót lại khẽ rỉ từ đầu khấc rơi lách tách xuống sàn đá bóng loáng."
    )
    add_body_paragraph(
        "Yên bừng tỉnh, khuôn mặt đỏ bừng đến tận mang tai, tim đập thình thịch như muốn nhảy khỏi lồng ngực. Cô hốt hoảng quay phắt người tháo chạy thục mạng ra phía ngoài."
    )
    add_body_paragraph(
        "\"Có muốn nếm thử một chút không?\" Tiếng cười khẩy đầy ngạo nghễ của Dương đuổi theo sau lưng cô, vọng qua dãy hành lang vắng. \"Anh vẫn để dành riêng cho em đấy!\""
    )
    add_body_paragraph(
        "Yên cắm đầu chạy như điên dại về phía thang máy, hơi thở dồn dập đứt quãng, nơi đáy chậu đã ướt sũng, dâm thủy tràn trề thấm đẫm cả chiếc quần lót mỏng manh."
    )

    add_ornament_divider("❦  ✦  ❦")

    # ==================== SCENE 5: TỰ THỎA MÃN CUỒNG LOẠN & CƠN NGHIỆN KHÔNG LỐI THOÁT ====================
    add_body_paragraph(
        "Vừa lảo đảo bước vào phòng khách sạn, Yên lập tức khóa chặt then cửa, ngã quỵ xuống tấm thảm trải sàn thở dốc hổn hển.",
        is_first=True
    )
    add_body_paragraph(
        "Ngọn lửa thèm khát và ngứa ngáy nơi vùng kín đã thiêu rụi toàn bộ lý trí còn sót lại của người đàn bà đức hạnh. Cô điên cuồng lột sạch toàn bộ váy áo trên người vứt tứ tung, trần như nhộng leo lên chiếc giường king-size rộng lớn phủ ga lụa đỏ thẫm, hai chân dạng rộng hết cỡ phơi bày đóa hoa xuân đang co thắt thèm khát."
    )
    add_body_paragraph(
        "\"Ưm... Dương... Dương ơi...\""
    )
    add_body_paragraph(
        "Tiếng rên rỉ ngọt ngào mà đê mê bật ra khỏi bờ môi khô rạo rực. Bàn tay cô run rẩy đưa xuống, ngón tay đè mạnh lên hạt le đang sưng đỏ, xoa miết cuồng nhiệt."
    )
    add_body_paragraph(
        "Trong đầu cô lúc này chỉ toàn là hình ảnh của Dương — tưởng tượng cây côn thịt gân guốc thô bạo đâm lút cán vào tận đáy tử cung mình, tưởng tượng dòng nước nóng bỏng mặn nồng phun ngập khoang miệng, tưởng tượng bản thân đang biến thành con đĩ dâm đãng nhất ngoan ngoãn nằm rên rỉ dưới thân hắn."
    )
    add_body_paragraph(
        "\"Ahh... Địt em đi... Thúc mạnh vào đi anh...\" Yên nức nở rên xiết, không còn chút liêm sỉ hay kiềm chế nào nữa."
    )
    add_body_paragraph(
        "Cô thọc hai ngón tay sâu vào lỗ lồn ướt nhẹp, móc ngoáy điên cuồng, tiếng nước nhóp nhép \"chẹp... chẹp... bạch... bạch...\" dâm loạn vang dội khắp căn phòng kín. Bàn tay còn lại vần vò hạt le và xoa bóp bầu ngực, đầu nhũ hoa cương cứng như hai hòn đá nhỏ."
    )
    add_body_paragraph(
        "\"Dương... Dương... Em ra... Em bắn mất rồi...\""
    )
    add_body_paragraph(
        "Yên khóc thét lên trong cơn khoái lạc tột cùng, hai bờ mông cong nảy bần bật trên nệm, dâm thủy ào ạt phun trào ướt đẫm cả bàn tay và ga giường. Toàn thân cô co giật từng hồi dữ dội trong cơn cực khoái muộn mằn, đôi môi vẫn run rẩy thì thào gọi mãi cái tên Dương."
    )
    add_body_paragraph(
        "Cùng lúc đó tại nhà, Khải vẫn đang say ngủ ngáy o o, hoàn toàn không hay biết người vợ sắp cưới dịu hiền của mình vừa trải qua một cơn thủ dâm điên dại, đắm chìm trong khoái cảm tội lỗi chỉ vì một gã đàn ông khác."
    )
    add_body_paragraph(
        "Kể từ đêm định mệnh ấy — đêm nào Yên cũng tự thỏa mãn chính mình. Đêm nào bờ môi cô cũng nức nở gọi tên Dương. Đêm nào cô cũng hoang tưởng về cây cặc gân guốc của hắn đang cuồng bạo giày xéo thân xác mình."
    )
    add_body_paragraph(
        "Yên đã thực sự nghiện Dương. Một cơn nghiện ngập đớn đau, tội lỗi ngập tràn nhưng vĩnh viễn không còn lối thoát."
    )

    # ==================== EMBEDDED ILLUSTRATION (IF PROVIDED) ====================
    if img_path and os.path.exists(img_path):
        table_ill = doc.add_table(rows=1, cols=1)
        table_ill.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_ill.autofit = False
        table_ill.columns[0].width = Inches(5.6)

        cell_ill = table_ill.cell(0, 0)
        set_cell_shading(cell_ill, HEX_BG_BOX)
        set_cell_box_border(cell_ill, HEX_BORDER_ROSE, size_eighth_pts=12)

        # Caption Tag
        cp1 = cell_ill.paragraphs[0]
        cp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp1.paragraph_format.space_before = Pt(8)
        cp1.paragraph_format.space_after = Pt(6)
        r_ill_tag = cp1.add_run("✦  M I N H   H Ọ A   C H Ư Ơ N G   4  ✦")
        r_ill_tag.font.name = "Palatino Linotype"
        r_ill_tag.font.size = Pt(9.5)
        r_ill_tag.font.bold = True
        r_ill_tag.font.color.rgb = COLOR_BURGUNDY

        # Image
        cp_img = cell_ill.add_paragraph()
        cp_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp_img.paragraph_format.space_before = Pt(2)
        cp_img.paragraph_format.space_after = Pt(6)
        run_img = cp_img.add_run()
        run_img.add_picture(img_path, width=Inches(4.8))

        # Caption text
        cp2 = cell_ill.add_paragraph()
        cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp2.paragraph_format.space_before = Pt(4)
        cp2.paragraph_format.space_after = Pt(8)
        r_ill_text = cp2.add_run("“Nằm trơ trọi giữa sắc đỏ nhung lụa, thể xác cồn cào ngứa ngáy và cơn mê muội tội lỗi gọi tên kẻ săn mồi”")
        r_ill_text.font.name = "Palatino Linotype"
        r_ill_text.font.size = Pt(11)
        r_ill_text.font.italic = True
        r_ill_text.font.bold = True
        r_ill_text.font.color.rgb = COLOR_WINE_DEEP

    # ==================== CHAPTER QUOTE BOX ====================
    table_q = doc.add_table(rows=1, cols=1)
    table_q.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_q.autofit = False
    table_q.columns[0].width = Inches(5.8)

    cell_q = table_q.cell(0, 0)
    set_cell_shading(cell_q, HEX_BG_BOX)
    set_cell_box_border(cell_q, HEX_BORDER_ROSE, size_eighth_pts=12)

    # Caption Tag
    cp1_q = cell_q.paragraphs[0]
    cp1_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp1_q.paragraph_format.space_before = Pt(10)
    cp1_q.paragraph_format.space_after = Pt(6)
    r_q_tag = cp1_q.add_run("✦  Đ I Ể M   N H Ấ N   C H Ư Ơ N G   4  ✦")
    r_q_tag.font.name = "Palatino Linotype"
    r_q_tag.font.size = Pt(9.5)
    r_q_tag.font.bold = True
    r_q_tag.font.color.rgb = COLOR_BURGUNDY

    # Caption text
    cp2_q = cell_q.add_paragraph()
    cp2_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp2_q.paragraph_format.space_before = Pt(4)
    cp2_q.paragraph_format.space_after = Pt(12)
    cp2_q.paragraph_format.left_indent = Inches(0.2)
    cp2_q.paragraph_format.right_indent = Inches(0.2)
    r_q_text = cp2_q.add_run("“Tội lỗi có thể đè nặng lên lý trí người đàn bà, nhưng cơn ngứa ngáy nơi xác thịt và nỗi ám ảnh về cây cặc của kẻ săn mồi... mới là thứ thực sự kéo cô vào vực thẳm sa đọa.”")
    r_q_text.font.name = "Palatino Linotype"
    r_q_text.font.size = Pt(11)
    r_q_text.font.italic = True
    r_q_text.font.bold = True
    r_q_text.font.color.rgb = COLOR_WINE_DEEP

    # ==================== END BLOCK ====================
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(24)
    p_end.paragraph_format.space_after = Pt(16)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("❖   HẾT CHƯƠNG 4   ❖")
    r_end.font.name = "Palatino Linotype"
    r_end.font.size = Pt(12.5)
    r_end.font.bold = True
    r_end.font.color.rgb = COLOR_BURGUNDY

    p_final = doc.add_paragraph()
    p_final.paragraph_format.space_before = Pt(4)
    p_final.paragraph_format.space_after = Pt(10)
    p_final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_final = p_final.add_run("❦  ❦  ❦")
    r_final.font.color.rgb = COLOR_ROSE_ACCENT
    r_final.font.size = Pt(11)

    doc.save(output_path)
    print(f"Document saved successfully at: {output_path}")

if __name__ == "__main__":
    img = r"C:\Users\Kien\.gemini\antigravity\scratch\truyen\public\images\illustrations\khat-tinh-ch4.jpg"
    out_download = r"C:\Users\Kien\Downloads\Khat_Tinh_Chuong_04.docx"
    out_scratch = r"C:\Users\Kien\.gemini\antigravity\scratch\truyen\Khat_Tinh_Chuong_04.docx"
    
    create_chapter_4_docx(out_download, img)
    create_chapter_4_docx(out_scratch, img)
